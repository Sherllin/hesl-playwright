"""Populate each downloaded template sheet; preserve header and workbook structure."""
import json
import sys
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.utils import get_column_letter

source, destination, sample_path = sys.argv[1:]
with open(source, "rb") as stream:
    book = load_workbook(stream)
with open(sample_path, encoding="utf-8") as stream:
    questions = json.load(stream)["questions"]
for sheet in book:
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)

def append(sheet, values):
    headers = [cell.value for cell in sheet[1]]
    assert all(key in headers for key in values), (sheet.title, values.keys(), headers)
    sheet.append([values.get(key) for key in headers])

for q in questions:
    if q['type'] not in book.sheetnames:
        # Custom names reuse the downloaded base-type columns, never invent headers.
        base_names = {1: '单选题', 2: '多选题', 3: '填空题', 4: '判断题', 5: '简答题', 6: '组合题', 7: '不定项选择题'}
        custom = book.copy_worksheet(book[base_names[q['typeId']]])
        custom.title = q['type']
    sheet = book[q["type"]]
    values = {"解析": q["analysis"], "难度": q["level"]}
    if 'knowledge' in q:
        values['知识点'] = '；'.join(q['knowledge'])
    if q["typeId"] == 6:
        append(sheet, {**values, "题目题干": q["stem"]})
        for child in q["children"]:
            answer = child.get('answer', '')
            if isinstance(answer, list):
                answer = '；'.join(f'第{i}空：{value}' for i, value in enumerate(answer, 1))
            append(sheet, {"小题题型": child["type"], "小题题干": child["stem"],
                           "选项内容": '；'.join(f'选项{chr(65+i)}：{value}' for i, value in enumerate(child.get('options', []))),
                           "正确答案": answer, "解析": child["analysis"], "难度": child["level"]})
        continue
    values["题干"] = q["stem"]
    if q["typeId"] == 3:
        values.update({f"答案{i}": answer for i, answer in enumerate(q["answer"], 1)})
    else:
        values["答案" if q["typeId"] == 5 else "正确答案"] = q["answer"]
    if "options" in q:
        values.update({"选项" + letter: value for letter, value in zip("ABCDEFGH", q["options"])})
    append(sheet, values)
    if q.get('media') == '嵌入图片' and q.get('image'):
        image = WorksheetImage(q['image'])
        image.height = round(180 * image.height / image.width)
        image.width = 180
        column = [cell.value for cell in sheet[1]].index('题干') + 1
        image.anchor = f'{get_column_letter(column)}{sheet.max_row}'
        sheet.add_image(image)
        sheet.row_dimensions[sheet.max_row].height = max(100, image.height * .75 + 24)
book.save(destination)
